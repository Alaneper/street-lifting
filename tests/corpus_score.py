"""Score de l'import sur le corpus synthétique (tests/gen_corpus.py).
Usage : python3 tests/corpus_score.py <dossier du corpus> [--detail FORMAT]
Score « exact » = exercice trouvé au bon jour avec le bon nom ET toutes ses séries (reps + charge) justes."""
import json, pathlib, subprocess, sys, collections
import openpyxl
H = pathlib.Path(__file__).resolve().parent; D = pathlib.Path(sys.argv[1])
idx = json.loads((D / "index.json").read_text())
def cell(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v)
inputs = []
for c in idx:
    if c["file"].endswith(".txt"): inputs.append({"id": c["id"], "text": (D / c["file"]).read_text()})
    else:
        wb = openpyxl.load_workbook(D / c["file"], data_only=True)
        inputs.append({"id": c["id"], "sheets": [{"name": ws.title, "rows": [[cell(v) for v in row] for row in ws.iter_rows(values_only=True)]} for ws in wb.worksheets]})
js = r"""const {parseProgram,parseWorkbook,fmtGroup}=require(process.argv[1]);let s='';process.stdin.on('data',d=>s+=d).on('end',()=>{
const out={};for(const c of JSON.parse(s)){const r=c.text!=null?parseProgram(c.text):parseWorkbook(c.sheets);
out[c.id]=r.weeks.map(w=>w.days.filter(d=>d.exos.length).map(d=>d.exos.map(e=>({canon:e.canon,name:e.name,
sets:e.groups.flatMap(g=>Array(Math.max(1,g.sets|0)).fill(String(g.reps)+'|'+(g.load?fmtGroup({sets:0,reps:0,load:g.load}).split(' · ')[1]:'')))}))))}
console.log(JSON.stringify(out))})"""
res = json.loads(subprocess.run(["node", "-e", js, str(H.parent / "src/prog_import.js")], input=json.dumps(inputs), capture_output=True, text=True, check=True).stdout)
detail = sys.argv[sys.argv.index("--detail") + 1] if "--detail" in sys.argv else None
agg = collections.defaultdict(lambda: [0, 0, 0, 0, 0]); NM = {"connus": [0, 0], "inédits": [0, 0]}   # exact, noms, total, structure ok, cas
for c in idx:
    P = res[c["id"]]; T = c["truth"]; a = agg[c["fmt"]]; a[4] += 1
    if len(P) >= len(T) and all(len(P[w]) >= len(T[w]) for w in range(len(T))): a[3] += 1
    misses = []
    for w, days in enumerate(T):
        for d, items in enumerate(days):
            pool = P[w][d] if w < len(P) and d < len(P[w]) else []
            used = set()
            for it in items:
                a[2] += 1
                want = [f"{it['reps']}|{it['load'] or ''}"] * it["sets"]
                j = next((j for j, e in enumerate(pool) if j not in used and e["canon"] == it["canon"]), None)
                if j is not None:
                    used.add(j); a[1] += 1
                    if pool[j]["sets"] == want: a[0] += 1
                    else: misses.append(f"S{w+1} J{d+1} {it['canon']}: attendu {want[0]}×{len(want)}, lu {pool[j]['sets'][:3]}{'…' if len(pool[j]['sets'])>3 else ''}")
                else: misses.append(f"S{w+1} J{d+1} {it['canon']} absent (lu : {[e['name'] for e in pool][:6]})")
    nk = "inédits" if c["test_names"] else "connus"; n_items = sum(len(i) for d in T for i in d)
    NM[nk][0] += n_items - len(misses); NM[nk][1] += n_items
    if detail and c["fmt"] == detail and misses: print(f"--- {c['id']} ({c['style']}{', noms inédits' if c['test_names'] else ''})"); print("   " + "\n   ".join(misses[:8]))
tot = [0, 0, 0]
print(f"{'format':22} {'exact':>7} {'noms':>7} {'structure':>10}")
for f, (e, n, t, s, k) in agg.items():
    tot[0] += e; tot[1] += n; tot[2] += t
    print(f"{f:22} {100*e/t:6.0f}% {100*n/t:6.0f}% {s:>5}/{k}")
tr = [sum(v[i] for f, v in agg.items() if not f.startswith("TEST")) for i in range(3)]
te = [sum(v[i] for f, v in agg.items() if f.startswith("TEST")) for i in range(3)]
print(f"\nFORMATS D'ENTRAÎNEMENT : exact {100*tr[0]/tr[2]:.0f} %  ·  noms {100*tr[1]/tr[2]:.0f} %")
if te[2]: print(f"FORMATS TEST (jamais réglés) : exact {100*te[0]/te[2]:.0f} %  ·  noms {100*te[1]/te[2]:.0f} %")
print(f"NOMS D'EXERCICES connus : exact {100*NM['connus'][0]/max(1,NM['connus'][1]):.0f} %  ·  noms jamais vus : exact {100*NM['inédits'][0]/max(1,NM['inédits'][1]):.0f} %")
print(f"TOTAL : {tot[0]}/{tot[2]} exercices exacts ({100*tot[0]/tot[2]:.0f} %)")
