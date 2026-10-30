"""Corpus de progs sophistiquées (chantier progs, étape 2, 30/10/2026). Données 100 % inventées.
Une prog = un artefact (note / csv / .xlsx / Sheet simulé) + sa vérité terrain.
Lecture fidèle à l'app :
 - .xlsx  → valeurs BRUTES des cellules (comme readXlsx de src/app.html : 0,75 reste « 0.75 », cellules fusionnées vides sauf la 1re, dates en nombre).
 - Sheet  → valeurs FORMATÉES telles que l'API Sheets les renvoie (« 75% », « 62,5 » en locale française, « 3:00 »).
 - note / csv → texte collé tel quel.
Vérité terrain : ce qu'un lecteur IDÉAL afficherait dans l'aperçu (fmtGroup du lecteur), pas ce que le lecteur actuel fait."""
import json, hashlib, pathlib, re, zipfile, xml.etree.ElementTree as ET

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "out"

def g(sets, reps, load=None, amrap=False):
    """Groupe de séries attendu. load = « 80 kg » | « 75 % » | « +20 kg » | « 225 lb » | « élastique » | None."""
    return (sets, str(reps), load, amrap)

def sig(grp):
    sets, reps, load, amrap = grp
    s = f"{sets}×{reps}" + ("+" if amrap and reps != "max" else "")
    return s + (" · " + load if load else "")

class Prog:
    def __init__(self, id, lot, kind, title, pieges, source):
        self.id, self.lot, self.kind, self.title, self.pieges, self.source = id, lot, kind, title, pieges, source
        self.truth = []        # {w,d,name,canon,groups,rpe,tempo,side}
        self.noise = []        # noms (normalisés) qui NE doivent PAS devenir des exercices (aide, réglages, totaux…)
        self.text = None       # note / csv
        self.sheets = None     # [{name, rows}] pour .xlsx (cellules brutes) ou Sheet simulé (valeurs formatées)
        self.xlsx = None       # [{name, rows, merges, fmt}] pour écrire le vrai .xlsx
        self.note = ""         # ce que cette prog cherche à casser
    def ex(self, w, d, name, groups, canon=None, rpe=None, tempo=None, side=False):
        self.truth.append(dict(w=w, d=d, name=name, canon=canon, groups=[list(x) for x in groups], rpe=rpe, tempo=tempo, side=side))
        return self
    def meta(self):
        return dict(id=self.id, lot=self.lot, kind=self.kind, title=self.title, pieges=self.pieges, source=self.source, note=self.note,
                    truth=self.truth, noise=self.noise)

def num(v):
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return "" if v is None else str(v)

# ---------- écriture .xlsx (openpyxl) ----------
def write_xlsx(p, path):
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for sh in p.xlsx:
        ws = wb.create_sheet(sh["name"])
        for r, row in enumerate(sh["rows"], 1):
            for c, v in enumerate(row, 1):
                if v is None or v == "": continue
                cell = ws.cell(row=r, column=c, value=v)
                f = sh.get("fmt", {}).get((r, c))
                if f: cell.number_format = f
                col = sh.get("fill", {}).get((r, c))
                if col: cell.fill = PatternFill("solid", start_color=col, end_color=col)
        for m in sh.get("merges", []): ws.merge_cells(m)
        if sh.get("hidden_rows"):
            for r in sh["hidden_rows"]: ws.row_dimensions[r].hidden = True
    wb.save(path)

# ---------- lecture .xlsx comme l'app (port de readXlsx, src/app.html ~3690) ----------
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
def read_xlsx_like_app(path):
    z = zipfile.ZipFile(path)
    def xml(n):
        try: return ET.fromstring(z.read(n))
        except KeyError: return None
    wb = xml("xl/workbook.xml"); rels = xml("xl/_rels/workbook.xml.rels")
    relT = {r.get("Id"): r.get("Target") for r in rels} if rels is not None else {}
    ss = xml("xl/sharedStrings.xml")
    shared = ["".join(t.text or "" for t in si.iter("{%s}t" % NS["m"])) for si in ss.findall("m:si", NS)] if ss is not None else []
    def colN(ref):
        n = 0
        for ch in re.sub(r"\d+", "", ref): n = n * 26 + ord(ch) - 64
        return n - 1
    out = []
    for sh in wb.find("m:sheets", NS):
        rid = sh.get("{%s}id" % NS["r"]); t = relT.get(rid)
        if not t: continue
        doc = xml("xl/" + re.sub(r"^/?(xl/)?", "", t))
        if doc is None: continue
        rows = []
        for row in doc.iter("{%s}row" % NS["m"]):
            r = (int(row.get("r") or len(rows) + 1)) - 1
            while len(rows) <= r: rows.append(None)
            rows[r] = rows[r] or []
            for c in row.findall("m:c", NS):
                ref, ty = c.get("r"), c.get("t"); v = c.find("m:v", NS); is_ = c.find("m:is", NS)
                if ty == "s" and v is not None: val = shared[int(v.text)] if int(v.text) < len(shared) else ""
                elif ty == "inlineStr" and is_ is not None: val = "".join(x.text or "" for x in is_.iter("{%s}t" % NS["m"]))
                elif v is not None: val = v.text or ""
                else: val = ""
                if ty == "b": val = "VRAI" if v is not None and v.text == "1" else "FAUX"
                k = colN(ref) if ref else len(rows[r])
                while len(rows[r]) <= k: rows[r].append("")
                rows[r][k] = re.sub(r"^(-?\d+)\.0$", r"\1", str(val))
        rows = [x or [] for x in rows]
        out.append({"name": sh.get("name"), "rows": rows})
    return out

# ---------- construction / empreinte ----------
def canonical(p):
    body = dict(kind=p.kind, text=p.text, sheets=p.sheets, truth=p.truth, noise=p.noise)
    return json.dumps(body, ensure_ascii=False, sort_keys=True)

def build(progs, lot):
    d = OUT / lot; d.mkdir(parents=True, exist_ok=True)
    index = []
    for p in progs:
        fn = None
        if p.kind in ("note", "csv"):
            fn = f"{p.id}.{'txt' if p.kind == 'note' else 'csv'}"; (d / fn).write_text(p.text, encoding="utf-8")
        elif p.kind == "xlsx":
            fn = f"{p.id}.xlsx"; write_xlsx(p, d / fn)
            p.sheets = read_xlsx_like_app(d / fn)     # ce que l'app verra vraiment
        elif p.kind == "sheet":
            fn = f"{p.id}.sheet.json"; (d / fn).write_text(json.dumps(p.sheets, ensure_ascii=False, indent=1), encoding="utf-8")
        m = p.meta(); m["file"] = fn; m["sha"] = hashlib.sha256(canonical(p).encode()).hexdigest()[:16]
        index.append(m)
    (d / "index.json").write_text(json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    return index

# ---------- aides d'écriture ----------
class TB:
    """Constructeur de note : écrit la ligne ET enregistre la vérité (semaine / jour courants)."""
    def __init__(self, p): self.p, self.L, self.w, self.d, self.wk0 = p, [], 0, -1, False
    def raw(self, s=""): self.L.append(s); return self
    def week(self, s):
        if self.wk0: self.w += 1
        self.wk0 = True; self.d = -1; return self.raw(s)
    def day(self, s): self.d += 1; return self.raw(s)
    def exo(self, line, name, groups, **kw): self.L.append(line); self.p.ex(self.w, max(self.d, 0), name, groups, **kw); return self
    def done(self): self.p.text = "\n".join(self.L) + "\n"; return self.p

def kg(v): return f"{float(str(v).replace(',', '.')):g} kg"
def lb(v): return f"{v} lb"
def pct(v): return f"{v} %"
def lest(v): return f"+{v} kg"
